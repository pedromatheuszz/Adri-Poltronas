"""
Restauracao de backup e neutralizacao de formula nas exportacoes.

RESTAURACAO
As copias eram criadas -- antes de cada migration e uma por dia -- e nao havia
caminho de volta: recuperar exigia fechar o app, achar a pasta em %APPDATA%,
entender o nome dos arquivos e sobrescrever o banco na mao. Um backup sem
procedimento de restauracao e uma suposicao, nao uma garantia; ninguem descobre
que a copia nao presta antes do dia em que precisa dela.

O download manual tinha um defeito proprio: enviava o ARQUIVO EM USO. Com o app
aberto isso captura um estado no meio de uma transacao e deixa de fora o que
estiver no `-wal` -- o operador guardava algo que parecia backup e podia nao
abrir.

PLANILHAS
Os dados exportados vem do cadastro, ou seja, de quem digitou. Um campo de
observacoes comecando com "=" vira formula ativa na planilha que o contador
abre. O nome do ataque e "CSV injection", mas ele nao e do CSV: e de quem abre
o arquivo -- por isso a defesa vale igual para XLSX.
"""
import io
import json
import os
import sqlite3
import sys
import tempfile
from contextlib import closing

import pytest

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
sys.path.insert(0, os.path.join(RAIZ, 'backend'))

import database  # noqa: E402
import planilha  # noqa: E402
import app as app_module  # noqa: E402
from erp.core import backup  # noqa: E402

ORIGEM_LOCAL = 'http://localhost'


@pytest.fixture
def client(tmp_path):
    caminho = str(tmp_path / "adri_poltronas.db")
    original = database.DB_PATH
    database.DB_PATH = caminho
    app_module.app.config['TESTING'] = True
    monkey_original = database.obter_caminho_banco_absoluto
    database.obter_caminho_banco_absoluto = lambda: caminho
    with app_module.app.app_context():
        database.criar_tabela()

    yield app_module.app.test_client()

    database.DB_PATH = original
    database.obter_caminho_banco_absoluto = monkey_original


def _criar_cliente(client, nome):
    return client.post('/api/clientes', headers={'Origin': ORIGEM_LOCAL},
                       data=json.dumps({
                           "nome": nome, "cep": "12345-678", "rua": "Rua Teste",
                           "bairro": "Centro", "cidade": "Sao Paulo",
                           "entrega_retirada": "Entrega",
                       }), content_type='application/json')


# =============================================================================
# Restauracao
# =============================================================================
def test_restaurar_devolve_o_banco_ao_estado_da_copia(client):
    """O teste que faltava. Um backup nunca exercitado e um arquivo, nao uma
    garantia."""
    _criar_cliente(client, 'Antes da copia')
    copia = backup.criar_backup(database.obter_caminho_banco_absoluto(), backup.MOTIVO_DIARIO)
    assert copia

    _criar_cliente(client, 'Depois da copia')
    assert len(json.loads(client.get('/api/clientes').data)) == 2

    resposta = client.post('/api/backup/restaurar', headers={'Origin': ORIGEM_LOCAL},
                           data=json.dumps({"arquivo": os.path.basename(copia)}),
                           content_type='application/json')

    assert resposta.status_code == 200
    assert json.loads(resposta.data)['requer_reinicio'] is True
    nomes = [c['nome'] for c in json.loads(client.get('/api/clientes').data)]
    assert nomes == ['Antes da copia']


def test_restaurar_guarda_o_estado_atual_antes_de_trocar(client):
    """Restaurar tambem e destrutivo: escolher o arquivo errado precisa ter
    volta."""
    caminho = database.obter_caminho_banco_absoluto()
    _criar_cliente(client, 'Original')
    copia = backup.criar_backup(caminho, backup.MOTIVO_DIARIO)
    _criar_cliente(client, 'Perdido na restauracao')

    antes = len(backup.listar_copias(caminho))
    client.post('/api/backup/restaurar', headers={'Origin': ORIGEM_LOCAL},
                data=json.dumps({"arquivo": os.path.basename(copia)}),
                content_type='application/json')

    copias = backup.listar_copias(caminho)
    assert len(copias) > antes
    salvaguarda = [c for c in copias if c['motivo'] == backup.MOTIVO_PRE_MIGRACAO]
    assert salvaguarda, "o estado anterior a restauracao nao foi preservado"

    # E a salvaguarda contem mesmo o que estava la.
    caminho_salvaguarda = os.path.join(
        backup.pasta_de_backups(caminho), salvaguarda[0]['arquivo']
    )
    with closing(sqlite3.connect(caminho_salvaguarda)) as conexao:
        nomes = {linha[0] for linha in conexao.execute("SELECT nome FROM clientes")}
    assert 'Perdido na restauracao' in nomes


def test_copia_corrompida_nao_e_restaurada(client):
    """Trocar um banco ruim por um pior seria o unico desfecho."""
    caminho = database.obter_caminho_banco_absoluto()
    _criar_cliente(client, 'Preservado')

    pasta = backup.pasta_de_backups(caminho)
    os.makedirs(pasta, exist_ok=True)
    ruim = os.path.join(pasta, 'adri_poltronas_diario_20200101_000000.db')
    with open(ruim, 'wb') as arquivo:
        arquivo.write(b'isto nao e um banco' * 40)

    resposta = client.post('/api/backup/restaurar', headers={'Origin': ORIGEM_LOCAL},
                           data=json.dumps({"arquivo": os.path.basename(ruim)}),
                           content_type='application/json')

    assert resposta.status_code == 500
    assert 'corrompida' in json.loads(resposta.data)['erro']
    nomes = [c['nome'] for c in json.loads(client.get('/api/clientes').data)]
    assert nomes == ['Preservado']


@pytest.mark.parametrize("nome_malicioso", [
    "../../Windows/system32/config/SAM",
    "..\\..\\etc\\passwd",
    "/etc/passwd",
    "subpasta/adri_poltronas_diario_20200101_000000.db",
])
def test_nome_de_arquivo_com_travessia_e_recusado(client, nome_malicioso):
    """O nome chega pela API. Sem a checagem, restaurar viraria leitura de
    arquivo arbitrario."""
    resposta = client.post('/api/backup/restaurar', headers={'Origin': ORIGEM_LOCAL},
                           data=json.dumps({"arquivo": nome_malicioso}),
                           content_type='application/json')
    assert resposta.status_code == 404


def test_listagem_traz_a_mais_recente_primeiro(client):
    caminho = database.obter_caminho_banco_absoluto()
    _criar_cliente(client, 'Um')
    pasta = backup.pasta_de_backups(caminho)
    os.makedirs(pasta, exist_ok=True)

    # Nomes controlados: o carimbo esta no proprio nome e e ele que ordena.
    for carimbo in ('20200101_000000', '20260101_000000'):
        origem = backup.criar_backup(caminho, backup.MOTIVO_DIARIO)
        os.replace(origem, os.path.join(pasta, f'adri_poltronas_diario_{carimbo}.db'))

    copias = json.loads(client.get('/api/backup/copias').data)['copias']
    assert copias[0]['arquivo'].startswith('adri_poltronas_diario_2026')


# =============================================================================
# Download manual: copia integra, nao o arquivo em uso
# =============================================================================
def test_download_do_banco_entrega_copia_que_abre(client):
    _criar_cliente(client, 'No download')

    resposta = client.get('/api/backup/sqlite')

    assert resposta.status_code == 200
    with tempfile.NamedTemporaryFile(suffix='.db', delete=False) as arquivo:
        arquivo.write(resposta.data)
        baixado = arquivo.name
    try:
        # `with sqlite3.connect(...)` fecha a TRANSACAO, nao a conexao -- e no
        # Windows o arquivo continua bloqueado. `closing` e o que solta.
        with closing(sqlite3.connect(baixado)) as conexao:
            assert conexao.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            nomes = {linha[0] for linha in conexao.execute("SELECT nome FROM clientes")}
        assert 'No download' in nomes
    finally:
        os.unlink(baixado)


def test_copia_temporaria_do_download_nao_fica_no_disco(client):
    """A copia intermediaria e o banco inteiro na pasta temporaria do sistema.
    Ela nao pode sobreviver ao download."""
    _criar_cliente(client, 'Temporario')
    antes = set(os.listdir(tempfile.gettempdir()))

    client.get('/api/backup/sqlite')

    novos = set(os.listdir(tempfile.gettempdir())) - antes
    assert not [n for n in novos if n.startswith('adri_poltronas_download_')]


def test_rota_antiga_de_backup_continua_funcionando(client):
    """Ela esta em instalacoes ja distribuidas e no atalho de teclado."""
    _criar_cliente(client, 'Compatibilidade')
    resposta = client.get('/api/clientes/backup/sqlite')
    assert resposta.status_code == 200
    assert resposta.data[:15] == b'SQLite format 3'


# =============================================================================
# Formula em planilha
# =============================================================================
@pytest.mark.parametrize("perigoso", [
    '=HYPERLINK("http://mal.exemplo","clique")',
    '+1+cmd|calc',
    '-2+3',
    '@SUM(A1:A9)',
    '\tvalor',
    '\rvalor',
])
def test_texto_que_viraria_formula_e_neutralizado(perigoso):
    resultado = planilha.neutralizar_formula(perigoso)
    assert resultado.startswith("'")
    assert resultado[1:] == perigoso


@pytest.mark.parametrize("inofensivo", ['Rua das Flores', '', 'Cliente 123'])
def test_texto_comum_passa_intacto(inofensivo):
    assert planilha.neutralizar_formula(inofensivo) == inofensivo


def test_numero_nao_vira_texto():
    """Prefixar numero quebraria toda soma na planilha do contador."""
    assert planilha.neutralizar_formula(1500) == 1500
    assert planilha.neutralizar_formula(-1500.5) == -1500.5
    assert planilha.neutralizar_formula(None) is None


def test_exportacao_csv_neutraliza_o_que_veio_do_cadastro(client):
    client.post('/api/clientes', headers={'Origin': ORIGEM_LOCAL},
                data=json.dumps({
                    "nome": "Cliente Normal", "cep": "12345-678", "rua": "Rua A",
                    "bairro": "Centro", "cidade": "Santos",
                    "entrega_retirada": "Entrega",
                    "observacoes": '=HYPERLINK("http://mal.exemplo","clique")',
                }), content_type='application/json')

    resposta = client.get('/api/relatorios/exportar?formato=csv')

    assert resposta.status_code == 200
    texto = resposta.data.decode('utf-8-sig')
    assert '"\'=HYPERLINK' in texto or "'=HYPERLINK" in texto
    assert ',=HYPERLINK' not in texto


def test_exportacao_xlsx_neutraliza_tambem(client):
    """O ataque nao e do CSV: e de quem abre o arquivo."""
    import openpyxl

    client.post('/api/clientes', headers={'Origin': ORIGEM_LOCAL},
                data=json.dumps({
                    "nome": "Cliente XLSX", "cep": "12345-678", "rua": "Rua A",
                    "bairro": "Centro", "cidade": "Santos",
                    "entrega_retirada": "Entrega",
                    "observacoes": '=1+1',
                }), content_type='application/json')

    resposta = client.get('/api/relatorios/exportar?formato=xlsx')
    assert resposta.status_code == 200

    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as arquivo:
        arquivo.write(resposta.data)
        caminho = arquivo.name
    try:
        planilha_lida = openpyxl.load_workbook(caminho)
        valores = [
            celula.value
            for linha in planilha_lida.active.iter_rows()
            for celula in linha
            if isinstance(celula.value, str) and '1+1' in celula.value
        ]
        assert valores == ["'=1+1"]
    finally:
        os.unlink(caminho)


# =============================================================================
# .xls: a interface prometia um suporte que nao existe
# =============================================================================
def test_xls_e_recusado_com_instrucao(client):
    """O leitor da importacao e o openpyxl, que so abre OOXML. Antes a extensao
    era aceita, o arquivo ia direto para `load_workbook` e o usuario recebia
    'nao foi possivel ler o arquivo' sem saber o que fazer."""
    # Com conteudo: um arquivo vazio para na checagem de tamanho antes de
    # chegar a de extensao, e o teste passaria a medir outra coisa.
    dados = {'arquivo': (io.BytesIO(b'\xd0\xcf\x11\xe0conteudo BIFF'), 'planilha.xls')}
    resposta = client.post('/api/clientes/importar', headers={'Origin': ORIGEM_LOCAL},
                           data=dados, content_type='multipart/form-data')

    assert resposta.status_code == 400
    mensagem = json.loads(resposta.data)['erro']
    assert '.xlsx' in mensagem or 'xls' in mensagem.lower()
