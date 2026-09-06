"""
Testes de regressão para os defeitos encontrados na auditoria técnica.

Cada teste referencia o item do relatório que o originou. Todos falham
na versão anterior do código.
"""
import os
import sys
import json
import tempfile

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'backend'))

import database  # noqa: E402
import planilha  # noqa: E402
import app as app_module  # noqa: E402

# O test client do Flask atende como host 'localhost' (sem porta). A checagem
# de origem deriva do host da requisicao, entao a origem valida acompanha isso.
ORIGEM_LOCAL = 'http://localhost'


@pytest.fixture
def client():
    with tempfile.NamedTemporaryFile(suffix='.db', delete=False) as f:
        db_path = f.name

    original = database.DB_PATH
    database.DB_PATH = db_path
    app_module.app.config['TESTING'] = True
    with app_module.app.app_context():
        database.criar_tabela()

    yield app_module.app.test_client()

    database.DB_PATH = original
    os.unlink(db_path)


def _criar_cliente_com_entrega(client, data_entrega):
    """
    Cria um cliente e uma entrega agendada na data informada.

    Ate a 1.5.0 este ajudante criava uma COBRANCA, e os testes de notificacao
    se apoiavam nela. Removido o modulo financeiro, pagamento deixou de gerar
    notificacao -- entrega gera, e e o que sustenta a mesma cobertura sem
    depender do que saiu.
    """
    client.post('/api/clientes', headers={'Origin': ORIGEM_LOCAL},
        data=json.dumps({
            "nome": "Cliente Teste", "cep": "12345-678", "rua": "Rua Teste",
            "bairro": "Centro", "cidade": "São Paulo", "entrega_retirada": "Entrega",
        }), content_type='application/json')
    cliente_id = json.loads(client.get('/api/clientes').data)[0]['id']
    database.agendar_entrega(cliente_id, "Cliente Teste", data_entrega,
                             descricao="Poltrona")
    return cliente_id


# -----------------------------------------------------------------------------
# A5 - sqlite3.Row nao tem .get(), o que derrubava TODAS as notificacoes
# -----------------------------------------------------------------------------






# -----------------------------------------------------------------------------
# M1 - CSRF: POST cross-site em endpoints que nao leem corpo
# -----------------------------------------------------------------------------




def test_referer_externo_tambem_e_recusado(client):
    resposta = client.post('/api/auth/logout',
                           headers={'Referer': 'https://site-malicioso.example/pagina'})
    assert resposta.status_code == 403


def test_get_nao_e_afetado_pela_checagem_de_origem(client):
    resposta = client.get('/api/clientes', headers={'Origin': 'https://qualquer.example'})
    assert resposta.status_code == 200


def test_loopback_alias_e_aceito_no_mesmo_host(client):
    """127.0.0.1 e localhost sao a mesma maquina, mas origens distintas no navegador."""
    resposta = client.post('/api/auth/logout', headers={'Origin': 'http://127.0.0.1'})
    assert resposta.status_code == 200


def test_checagem_de_origem_acompanha_a_porta_do_servidor(client):
    """
    A lista de origens e derivada do host da requisicao, e nao fixada em :5000.
    Uma lista fixa faria o proprio frontend levar 403 em qualquer outra porta.
    """
    aceita = client.post('/api/auth/logout',
                         headers={'Origin': 'http://localhost:8080', 'Host': 'localhost:8080'})
    assert aceita.status_code == 200

    recusada = client.post('/api/auth/logout',
                           headers={'Origin': 'http://localhost:9999', 'Host': 'localhost:8080'})
    assert recusada.status_code == 403


def test_sem_origin_e_sem_referer_e_permitido(client):
    """Clientes locais (curl, scripts) nao sao navegadores e nao sofrem CSRF."""
    assert client.post('/api/auth/logout').status_code == 200


# -----------------------------------------------------------------------------
# A2 - webhook Asaas: token estatico `asaas-access-token`, nao HMAC
# -----------------------------------------------------------------------------














def test_host_externo_continua_recusado_no_resto_da_api(client):
    """A isencao vale para UMA rota, a que chega de fora por desenho. O vetor
    de DNS rebinding contra o resto da API segue fechado."""
    app_module.app.config['TESTING'] = False
    try:
        resposta = client.get('/api/clientes', headers={'Host': 'minhaloja.com.br'})
    finally:
        app_module.app.config['TESTING'] = True

    assert resposta.status_code == 400


# -----------------------------------------------------------------------------
# A1 - asaas lia os.environ no import, antes do load_dotenv
# -----------------------------------------------------------------------------








# -----------------------------------------------------------------------------
# M5 - cabecalhos de seguranca
# -----------------------------------------------------------------------------
def test_resposta_traz_cabecalhos_de_seguranca(client):
    cabecalhos = client.get('/api/clientes').headers

    assert "default-src 'self'" in cabecalhos['Content-Security-Policy']
    assert cabecalhos['X-Content-Type-Options'] == 'nosniff'
    assert cabecalhos['X-Frame-Options'] == 'DENY'


def test_csp_libera_apenas_as_origens_externas_realmente_usadas(client):
    csp = client.get('/').headers['Content-Security-Policy']

    assert 'https://viacep.com.br' in csp
    assert 'https://nominatim.openstreetmap.org' in csp
    assert "frame-ancestors 'none'" in csp


# -----------------------------------------------------------------------------
# B1 - planilha compartilhada entre exportacao e OneDrive
# -----------------------------------------------------------------------------
def test_planilha_escreve_cabecalhos_e_valores():
    import openpyxl

    buffer = planilha.construir_planilha(
        "Clientes", ["Nome", "Cidade"], ['nome', 'cidade'],
        [{'nome': 'Ana', 'cidade': 'Santos'}, {'nome': 'Bruno', 'cidade': 'Recife'}],
    )
    ws = openpyxl.load_workbook(buffer).active

    assert ws.title == "Clientes"
    assert [c.value for c in ws[1]] == ["Nome", "Cidade"]
    assert [c.value for c in ws[2]] == ["Ana", "Santos"]
    assert ws.max_row == 3


def test_planilha_tolera_registro_sem_o_campo():
    import openpyxl

    buffer = planilha.construir_planilha(
        "X", ["Nome", "Obs"], ['nome', 'observacoes'], [{'nome': 'Ana'}],
    )
    ws = openpyxl.load_workbook(buffer).active

    assert ws.cell(row=2, column=2).value in (None, '')


# -----------------------------------------------------------------------------
# M2 - Alembic resolvia alembic.ini pelo CWD e nunca rodava de fato
# -----------------------------------------------------------------------------


def test_migracao_inicial_usa_tipos_validos_do_sqlalchemy():
    """
    A migracao usava `sa.Real()`, que nao existe — erro invisivel enquanto o
    Alembic nunca chegava a executar.
    """
    import re
    import sqlalchemy as sa

    caminho = os.path.join(
        os.path.dirname(__file__), '..', 'backend', 'migrations', 'versions',
        '21d8d4eb2b6f_initial_schema.py')
    with open(caminho, encoding='utf-8') as arquivo:
        origem = arquivo.read()

    desconhecidos = [t for t in set(re.findall(r'sa\.([A-Za-z_]+)\(', origem))
                     if not hasattr(sa, t)]
    assert not desconhecidos, f"Tipos inexistentes no SQLAlchemy: {desconhecidos}"


# -----------------------------------------------------------------------------
# UI - dialogos nativos substituidos e pagina de callback tematizada
# -----------------------------------------------------------------------------
def _ler_app_js():
    caminho = os.path.join(os.path.dirname(__file__), '..', 'frontend', 'js', 'app.js')
    with open(caminho, encoding='utf-8') as arquivo:
        return arquivo.read()


def _ler_template_principal():
    caminho = os.path.join(os.path.dirname(__file__), '..', 'backend', 'templates', 'index.html')
    with open(caminho, encoding='utf-8') as arquivo:
        return arquivo.read()


def _ler_css_principal():
    caminho = os.path.join(os.path.dirname(__file__), '..', 'frontend', 'css', 'style.css')
    with open(caminho, encoding='utf-8') as arquivo:
        return arquivo.read()


def test_template_principal_usa_estrutura_html5_semantica():
    html = _ler_template_principal()

    assert html.lstrip().lower().startswith('<!doctype html>')
    for tag in ('header', 'nav', 'main', 'section', 'article', 'search', 'dialog', 'footer'):
        assert f'<{tag}' in html, f'Elemento HTML5 ausente: <{tag}>'

    # A garantia e que TODO modal seja <dialog>, e nao uma contagem fixa: o
    # numero cresce a cada tela nova, e travar nele so gera falso alarme. O
    # que importa e que ninguem volte a usar <div> -- foi de la que vinham o
    # foco passeando por tras do modal e o Esc que nao fechava.
    assert html.count('<dialog class="modal-overlay"') >= 5
    assert '<div class="modal-overlay"' not in html


def test_dialogos_html5_usam_api_nativa_e_backdrop_css():
    javascript = _ler_app_js()
    css = _ler_css_principal()

    assert '.showModal()' in javascript
    assert "document.querySelector('dialog[open]')" in javascript
    assert '.modal-overlay[open]' in css
    assert '.modal-overlay::backdrop' in css
    assert '[hidden] { display: none !important; }' in css


def test_frontend_nao_usa_dialogos_nativos():
    """
    confirm()/alert() nativos ignoram o tema e exibem o endereco do servidor
    ("127.0.0.1:5000 diz"). Devem passar por confirmarAcao().
    """
    import re

    codigo = _ler_app_js()
    sem_comentarios = re.sub(r'//.*', '', codigo)
    # O lookbehind ja descarta `window.confirm(`, mantido apenas como fallback
    # interno de confirmarAcao() para o caso do modal nao existir no DOM.
    chamadas = re.findall(r'(?<![.\w])(confirm|alert|prompt)\s*\(', sem_comentarios)
    assert chamadas == [], f"Dialogos nativos remanescentes: {chamadas}"


def test_confirmar_acao_esta_definida_e_modal_existe_no_template():
    assert 'function confirmarAcao(' in _ler_app_js()

    caminho = os.path.join(os.path.dirname(__file__), '..', 'backend', 'templates', 'index.html')
    with open(caminho, encoding='utf-8') as arquivo:
        html = arquivo.read()
    for elemento in ('id="modalConfirmar"', 'id="modalConfirmarTitulo"',
                     'id="modalConfirmarMsg"', 'id="modalConfirmarOk"',
                     'id="modalConfirmarCancelar"'):
        assert elemento in html, f"Faltando no template: {elemento}"


def test_callback_de_erro_renderiza_como_erro(client):
    """
    O template tem `{% if erro %}`, mas app.py nunca passava a variavel: toda
    tela de falha do login aparecia com o check verde de sucesso.
    """
    resposta = client.get('/callback')  # sem `code` -> erro

    assert resposta.status_code == 400
    corpo = resposta.data.decode('utf-8')
    assert 'var(--erro)' in corpo, "Tela de erro nao usa a cor de erro"
    assert 'var(--ok)' not in corpo, "Tela de erro herdou o visual de sucesso"


def test_callback_usa_a_paleta_do_app(client):
    """
    A tela de callback abre no navegador do sistema e nao alcanca
    frontend/css/tokens.css, entao a paleta e duplicada em config.py. Este
    teste e o que impede as duas de divergirem em silencio.

    Atualizado junto com a renovacao visual: o app deixou de ser azul-escuro
    e passou a ser claro e quente, e a tela de login seguiu junto. Antes daqui
    o teste fixava `#0f172a`, o fundo do tema antigo.
    """
    corpo = client.get('/callback').data.decode('utf-8')

    assert '#F2EEE7' in corpo, "Callback nao usa o fundo claro do app"
    assert '#0f172a' not in corpo, "Callback ainda tem o fundo azul-escuro antigo"
    assert 'background:#f4f6f9' not in corpo, "Callback ainda tem o fundo claro antigo"


# -----------------------------------------------------------------------------
# Notificacoes lidas
# -----------------------------------------------------------------------------


def test_marcar_e_desmarcar_notificacao_como_lida(client):
    _criar_cliente_com_entrega(client, '2020-01-01')
    chave = json.loads(client.get('/api/notificacoes').data)['notificacoes'][0]['chave']

    marcar = client.post('/api/notificacoes/lidas', headers={'Origin': ORIGEM_LOCAL},
                         data=json.dumps({'chaves': [chave]}), content_type='application/json')
    assert marcar.status_code == 200
    dados = json.loads(client.get('/api/notificacoes').data)
    assert dados['nao_lidas'] == 0
    assert all(n['lida'] for n in dados['notificacoes'])

    client.delete('/api/notificacoes/lidas', headers={'Origin': ORIGEM_LOCAL},
                  data=json.dumps({'chaves': [chave]}), content_type='application/json')
    assert json.loads(client.get('/api/notificacoes').data)['nao_lidas'] == 1


def test_marcar_como_lida_e_idempotente(client):
    _criar_cliente_com_entrega(client, '2020-01-01')
    chave = json.loads(client.get('/api/notificacoes').data)['notificacoes'][0]['chave']
    corpo = json.dumps({'chaves': [chave]})

    for _ in range(3):
        client.post('/api/notificacoes/lidas', headers={'Origin': ORIGEM_LOCAL},
                    data=corpo, content_type='application/json')

    assert json.loads(client.get('/api/notificacoes').data)['nao_lidas'] == 0


def test_marcar_todas_como_lidas(client):
    _criar_cliente_com_entrega(client, '2020-01-01')
    database.agendar_entrega(1, 'Cliente Teste', '2020-01-01', '', 'Atrasada')
    assert json.loads(client.get('/api/notificacoes').data)['nao_lidas'] >= 2

    client.post('/api/notificacoes/lidas', headers={'Origin': ORIGEM_LOCAL},
                data=json.dumps({'todas': True}), content_type='application/json')

    assert json.loads(client.get('/api/notificacoes').data)['nao_lidas'] == 0


def test_nao_lidas_aparecem_antes_das_lidas(client):
    _criar_cliente_com_entrega(client, '2020-01-01')
    database.agendar_entrega(1, 'Cliente Teste', '2020-01-01', '', 'Atrasada')
    notifs = json.loads(client.get('/api/notificacoes').data)['notificacoes']
    client.post('/api/notificacoes/lidas', headers={'Origin': ORIGEM_LOCAL},
                data=json.dumps({'chaves': [notifs[0]['chave']]}), content_type='application/json')

    ordem = [n['lida'] for n in json.loads(client.get('/api/notificacoes').data)['notificacoes']]

    assert ordem == sorted(ordem), f"Lidas deveriam vir por ultimo: {ordem}"


def test_chaves_invalidas_sao_ignoradas(client):
    resposta = client.post('/api/notificacoes/lidas', headers={'Origin': ORIGEM_LOCAL},
                           data=json.dumps({'chaves': ['', '   ']}), content_type='application/json')
    assert resposta.status_code == 200
    assert json.loads(resposta.data)['quantidade'] == 0


# -----------------------------------------------------------------------------
# Dashboard financeiro
# -----------------------------------------------------------------------------
def _receber(cliente_id, valor, data_pagamento, vencimento=None):
    pagamento_id = database.salvar_pagamento(
        cliente_id, 'Cliente Teste', valor, 'Locacao', vencimento or data_pagamento[:10])
    with database.conectar() as conexao:
        conexao.execute("UPDATE pagamentos SET status='Recebido', data_pagamento=? WHERE id=?",
                        (data_pagamento, pagamento_id))
    return pagamento_id














def test_dashboard_nao_exibe_mais_o_texto_atualizado():
    caminho = os.path.join(os.path.dirname(__file__), '..', 'backend', 'templates', 'index.html')
    with open(caminho, encoding='utf-8') as arquivo:
        html = arquivo.read()

    assert 'metricAtualizado' not in html
    assert 'Atualizado:' not in html
    # Os canvases precisam existir, senao o codigo dos graficos nunca roda.
    # `chartReceitaMensal` saiu na 1.5.0 com o modulo financeiro.
    for canvas in ('chartEntregaRetirada', 'chartCidades', 'chartTimeline'):
        assert f'id="{canvas}"' in html, f"Canvas ausente: {canvas}"
    assert 'chartReceitaMensal' not in html




# -----------------------------------------------------------------------------
# Varredura profunda - incoerencias de logica entre camadas
# -----------------------------------------------------------------------------








def test_entrega_cancelada_nao_gera_alerta(client):
    """
    As consultas de notificacao so excluiam 'Concluida', mas 'Cancelada' e um
    status valido em blueprints/entregas.py: a entrega cancelada continuava
    aparecendo como agendada e, passada a data, como ENTREGA ATRASADA.
    """
    cliente_id = _criar_cliente_com_entrega(client, '2030-01-01')
    entrega_id = database.agendar_entrega(cliente_id, 'Cliente Teste', '2020-01-01', '', 'Poltrona')

    assert any(n['tipo'] == 'entrega_atrasada' for n in database.obter_notificacoes())

    database.atualizar_status_entrega(entrega_id, 'Cancelada')

    assert not any(n['tipo'] == 'entrega_atrasada' for n in database.obter_notificacoes())


def test_erro_do_dashboard_nao_vaza_detalhe_interno(client, monkeypatch):
    """
    dashboard.py e notificacoes.py devolviam str(e) ao navegador, contra o
    handler central de app.py e o API_STANDARD.md ("sem expor o stack trace").
    """
    def explode():
        raise RuntimeError("caminho/interno/secreto.db travou")

    monkeypatch.setattr(database, 'obter_dados_grafico_por_tipo', explode)
    resposta = client.get('/api/dashboard/dados')

    assert resposta.status_code == 500
    corpo = json.loads(resposta.data)
    assert corpo['erro'] == "Erro interno do servidor."
    assert 'id' in corpo
    assert 'secreto' not in resposta.get_data(as_text=True)


# =============================================================================
# A migracao emudecia o log do aplicativo
#
# Achado que apareceu enquanto se testava outra coisa: os testes do log do
# Asaas passavam sozinhos e capturavam ZERO registro quando rodavam depois
# deste arquivo. A causa nao era do teste.
#
# `fileConfig` tem `disable_existing_loggers=True` por padrao, e o env.py do
# Alembic o chamava sem argumento. Todo logger ja criado e nao declarado em
# alembic.ini era desligado -- e `servidor.py` importa `app` (que cria
# `adri.app`, `adri.asaas`, `adri.security`, `adri.db`, `adri.capacidade`)
# antes de chamar `preparar_banco()`, que roda `command.upgrade` em todo boot.
#
# Resultado em producao: a partir da primeira migracao de cada execucao, o
# aplicativo parava de escrever no proprio app.log. Sem estardalhaco -- o
# arquivo continuava la, so nao recebia mais os avisos de AUTH, CAPACIDADE e
# SEGURANCA, nem o log que o QUICK_REFERENCE manda o operador enviar quando
# algo da errado.
# =============================================================================


def _migrar_em_banco_novo(tmp_path, monkeypatch):
    import migrations_manager

    monkeypatch.setenv('ADRI_DB_PATH', str(tmp_path / 'migrar.db'))
    monkeypatch.setattr(database, 'DB_PATH', str(tmp_path / 'migrar.db'))
    migrations_manager.aplicar_migrações()


# Loggers que o aplicativo cria e que a migracao nao pode emudecer.
# `adri.asaas` saiu da lista junto com a integracao.
NOMES_DE_LOGGER = ['adri.app', 'adri.security', 'adri.db', 'adri.capacidade',
                   'adri.licenca']


def test_a_migracao_nao_desliga_os_loggers_do_aplicativo(tmp_path, monkeypatch):
    """Metade do estrago: `disable_existing_loggers`, que e True por padrao."""
    import logging

    loggers = [logging.getLogger(nome) for nome in NOMES_DE_LOGGER]
    for logger in loggers:
        logger.disabled = False

    _migrar_em_banco_novo(tmp_path, monkeypatch)

    desligados = [logger.name for logger in loggers if logger.disabled]
    assert not desligados, f"a migracao emudeceu {desligados}"


def test_a_migracao_nao_leva_embora_o_arquivo_de_log(tmp_path, monkeypatch):
    """
    A outra metade, e a mais grave.

    `[logger_root]` do alembic.ini declara `handlers = console`, e `fileConfig`
    SUBSTITUI os handlers da raiz em vez de acrescentar. O RotatingFileHandler
    que escreve o `app.log` em %APPDATA% ia embora junto -- entao nem WARNING
    nem ERROR chegavam mais ao arquivo. Some a trilha de [AUTH], [CAPACIDADE]
    e [SEGURANCA], e some o log que o QUICK_REFERENCE manda o operador enviar.
    """
    import logging

    raiz = logging.getLogger()
    marcador = logging.NullHandler()
    raiz.addHandler(marcador)
    try:
        _migrar_em_banco_novo(tmp_path, monkeypatch)
        assert marcador in raiz.handlers,             "a migracao trocou os handlers da raiz pelos do alembic.ini"
    finally:
        if marcador in raiz.handlers:
            raiz.removeHandler(marcador)


def test_a_migracao_nao_rebaixa_o_nivel_da_raiz(tmp_path, monkeypatch):
    """`level = WARN` no alembic.ini derrubava todo `log.info` do aplicativo:
    os loggers `adri.*` sao NOTSET e herdam o nivel da raiz."""
    import logging

    raiz = logging.getLogger()
    nivel_antes = raiz.level
    raiz.setLevel(logging.INFO)
    try:
        _migrar_em_banco_novo(tmp_path, monkeypatch)
        assert raiz.level == logging.INFO,             f"a raiz foi rebaixada para {logging.getLevelName(raiz.level)}"
    finally:
        raiz.setLevel(nivel_antes)


def test_o_env_do_alembic_declara_a_intencao():
    """Guarda contra a regressao mais provavel: alguem regenerar o env.py a
    partir do modelo do Alembic, que vem sem o argumento."""
    caminho = os.path.join(os.path.dirname(__file__), '..', 'backend',
                           'migrations', 'env.py')
    with open(caminho, encoding='utf-8') as arquivo:
        fonte = arquivo.read()

    assert 'disable_existing_loggers=False' in fonte


# -----------------------------------------------------------------------------
# PERF-06 - `notificacoes_lidas` so crescia.
#
# Cada entrega e cada cliente que o operador dispensasse deixava uma linha para
# sempre -- inclusive depois de a entrega ser concluida ou o cliente inativado,
# quando a notificacao ja nem e gerada e a marca perdeu qualquer funcao.
#
# O criterio de expurgo e a existencia do FATO, e nao a idade da marca. Um
# prazo em dias traria o defeito oposto: uma entrega ainda atrasada e ainda
# dispensada voltaria a aparecer sozinha ao vencer o prazo, sem o operador
# entender por que. Estes testes fixam as duas metades.
# -----------------------------------------------------------------------------
def test_marca_de_lida_some_quando_a_entrega_e_concluida(client):
    cliente_id = _criar_cliente_com_entrega(client, '2020-01-05')
    with database.conectar() as conexao:
        entrega_id = conexao.execute(
            'SELECT id FROM entregas_agendadas WHERE cliente_id=?', (cliente_id,)
        ).fetchone()[0]

    database.marcar_notificacoes_lidas([f'entrega_atrasada:{entrega_id}'])
    database.atualizar_status_entrega(entrega_id, 'Concluida')

    database.obter_notificacoes()

    with database.conectar() as conexao:
        restantes = [linha[0] for linha in
                     conexao.execute('SELECT chave FROM notificacoes_lidas')]
    assert f'entrega_atrasada:{entrega_id}' not in restantes


def test_marca_de_lida_sobrevive_enquanto_a_notificacao_vale(client):
    """A metade que importa mais: dispensar nao pode ser desfeito sozinho."""
    cliente_id = _criar_cliente_com_entrega(client, '2020-01-05')
    with database.conectar() as conexao:
        entrega_id = conexao.execute(
            'SELECT id FROM entregas_agendadas WHERE cliente_id=?', (cliente_id,)
        ).fetchone()[0]

    chave = f'entrega_atrasada:{entrega_id}'
    database.marcar_notificacoes_lidas([chave])

    notificacoes = database.obter_notificacoes()

    atrasada = [n for n in notificacoes if n['chave'] == chave]
    assert atrasada, "a entrega continua atrasada; a notificacao tem de existir"
    assert atrasada[0]['lida'] is True, "a marca de lida foi expurgada indevidamente"


def test_marca_orfa_de_registro_inexistente_e_removida(client):
    """Chave apontando para id que nao existe mais nao tem como voltar a valer."""
    _criar_cliente_com_entrega(client, '2020-01-05')
    database.marcar_notificacoes_lidas(['entrega_hoje:99999', 'cliente_inativo:88888'])

    database.obter_notificacoes()

    with database.conectar() as conexao:
        restantes = [linha[0] for linha in
                     conexao.execute('SELECT chave FROM notificacoes_lidas')]
    assert 'entrega_hoje:99999' not in restantes
    assert 'cliente_inativo:88888' not in restantes
