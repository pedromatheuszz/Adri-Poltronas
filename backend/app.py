import logging
import os
import secrets
import hashlib
import base64
import uuid
import json
import sys
import time
import io
import csv
import re
import urllib.request
import urllib.parse
import urllib.error
import webbrowser
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import (
    Flask,
    request,
    jsonify,
    send_from_directory,
    render_template,
    render_template_string,
    send_file,
)
from werkzeug.exceptions import HTTPException

import database
import validators
import asaas

LOG_DIR = os.path.join(os.environ.get("APPDATA", os.getcwd()), "AdriPoltronas", "logs")
try:
    os.makedirs(LOG_DIR, exist_ok=True)
    _log_path = os.path.join(LOG_DIR, "app.log")
    _handlers = [logging.FileHandler(_log_path, encoding="utf-8"), logging.StreamHandler(sys.stderr)]
except OSError:
    _handlers = [logging.StreamHandler(sys.stderr)]

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    handlers=_handlers,
)
log = logging.getLogger("adri.app")

diretorio_atual = os.path.dirname(os.path.abspath(__file__))
if diretorio_atual not in sys.path:
    sys.path.append(diretorio_atual)

PASTA_FRONTEND = os.path.abspath(os.path.join(diretorio_atual, '..', 'frontend'))
app = Flask(__name__,
    static_folder=PASTA_FRONTEND,
    static_url_path='',
    template_folder=os.path.join(diretorio_atual, 'templates'),
)
app.secret_key = os.environ.get('ADRI_FLASK_SECRET')
if not app.secret_key:
    if getattr(sys, 'frozen', False):
        log.warning("ADRI_FLASK_SECRET nao definida. Usando chave temporaria.")
    app.secret_key = secrets.token_hex(32)

CLIENT_ID = os.environ.get('ADRI_MICROSOFT_CLIENT_ID', '14bc533a-a8a2-4e2b-9142-773dcd117269')
REDIRECT_URI = "http://localhost:5000/callback"
GEMINI_API_KEY = os.environ.get('GEMINI_API_KEY', '')
if not GEMINI_API_KEY:
    log.warning("GEMINI_API_KEY nao definida. Sugestao de bairro via IA desabilitada.")

ASAAS_API_KEY = os.environ.get('ASAAS_API_KEY', '')
if not ASAAS_API_KEY:
    log.warning("ASAAS_API_KEY nao definida. Pagamentos online desabilitados.")

_oauth_pending: dict = {}


def _gerar_code_verifier() -> str:
    return base64.urlsafe_b64encode(secrets.token_bytes(32)).rstrip(b'=').decode('ascii')


def _derivar_code_challenge(verifier: str) -> str:
    digest = hashlib.sha256(verifier.encode('ascii')).digest()
    return base64.urlsafe_b64encode(digest).rstrip(b'=').decode('ascii')


@app.errorhandler(HTTPException)
def _erro_http(e):
    if request.path.startswith('/api/'):
        return jsonify({"erro": e.description or e.name}), e.code
    return e


@app.errorhandler(Exception)
def _erro_generico(e):
    correlation_id = uuid.uuid4().hex[:12]
    log.exception("Erro nao tratado [%s]: %s", correlation_id, e)
    if request.path.startswith('/api/'):
        return jsonify({"erro": "Erro interno do servidor.", "id": correlation_id}), 500
    raise e


@app.route('/')
def index():
    m = database.obter_metricas_dashboard()
    total = m['total_clientes'] or 1
    pct = round(((m['cadastros_30d'] or 0) / total) * 100)
    return render_template('index.html',
        TOTAL=m['total_clientes'],
        ENTREGAS=m['total_entrega'],
        RETIRADAS=m['total_retirada'],
        NOVOS=pct,
        NOW=time.strftime('%d/%m/%Y %H:%M'),
    )


@app.route('/manifest.json')
def manifest():
    return send_from_directory(app.static_folder, 'manifest.json')


@app.route('/sw.js')
def service_worker():
    return send_from_directory(app.static_folder, 'sw.js')


# -----------------------------------------------------------------------------
# OAuth Microsoft
# -----------------------------------------------------------------------------
@app.route('/api/auth/login')
def auth_login():
    agora = time.time()
    chaves_antigas = [k for k, v in _oauth_pending.items() if agora - v.get('ts', 0) > 300]
    for k in chaves_antigas:
        _oauth_pending.pop(k, None)

    code_verifier = _gerar_code_verifier()
    code_challenge = _derivar_code_challenge(code_verifier)
    state = secrets.token_urlsafe(16)
    _oauth_pending[state] = {'verifier': code_verifier, 'ts': agora}
    log.info("[LOGIN] state=%s", state)

    scopes = "User.Read Files.ReadWrite"
    url = (
        f"https://login.microsoftonline.com/common/oauth2/v2.0/authorize"
        f"?client_id={CLIENT_ID}&response_type=code&redirect_uri={REDIRECT_URI}"
        f"&response_mode=query&scope={urllib.parse.quote(scopes)}"
        f"&code_challenge={code_challenge}&code_challenge_method=S256"
        f"&state={urllib.parse.quote(state)}"
    )
    webbrowser.open(url)
    return jsonify({"mensagem": "Navegador aberto"})


_CALLBACK_TEMPLATE = """\
<html>
<body style="font-family:sans-serif; text-align:center; padding-top:60px; background:#f4f6f9; color:#333;">
    <div style="max-width:480px; margin:0 auto; background:white; border-radius:12px; padding:40px 32px; box-shadow:0 8px 30px rgba(0,0,0,0.08);">
        <div style="font-size:48px; margin-bottom:16px;">{% if erro %}&#10060;{% else %}&#9989;{% endif %}</div>
        <h2 style="color:{% if erro %}#e74c3c{% else %}#27ae60{% endif %}; margin:0 0 12px;">{{ titulo }}</h2>
        <p style="color:#555; line-height:1.6; margin:0 0 8px;">{{ mensagem }}</p>
        <p style="color:#888; font-size:0.9rem; margin:0;">{{ instrucao }}</p>
        {% if not erro %}
        <p style="color:#aaa; font-size:0.8rem; margin-top:24px;">Esta aba pode ser fechada manualmente.</p>
        {% endif %}
    </div>
    <script>try { setTimeout(function(){ window.close(); }, 3000); } catch(e) {}</script>
</body>
</html>
"""


@app.route('/callback')
def auth_callback():
    code = request.args.get('code')
    state_recebido = request.args.get('state', '')
    log.info("[CALLBACK] state=%s, code=%s...", state_recebido, (code or '')[:8])

    if not code:
        return render_template_string(_CALLBACK_TEMPLATE, titulo="Erro na autenticacao",
            mensagem="Codigo de autorizacao ausente.", instrucao="Feche esta aba e tente novamente."), 400

    dados_pendentes = _oauth_pending.pop(state_recebido, None) if state_recebido else None
    if not dados_pendentes:
        log.warning("[CALLBACK] state '%s' not found.", state_recebido)
        return render_template_string(_CALLBACK_TEMPLATE, titulo="Erro de seguranca",
            mensagem="Sessao de login invalida ou expirada.", instrucao="Feche esta aba e reinicie o login."), 400

    code_verifier = dados_pendentes['verifier']
    url_token = "https://login.microsoftonline.com/common/oauth2/v2.0/token"
    payload = {
        'client_id': CLIENT_ID, 'scope': 'User.Read Files.ReadWrite', 'code': code,
        'redirect_uri': REDIRECT_URI, 'grant_type': 'authorization_code', 'code_verifier': code_verifier,
    }
    data = urllib.parse.urlencode(payload).encode('utf-8')

    try:
        with urllib.request.urlopen(urllib.request.Request(url_token, data=data, method='POST')) as resposta:
            res_dados = json.loads(resposta.read().decode('utf-8'))
            token_acesso = res_dados.get('access_token')

            req_perfil = urllib.request.Request("https://graph.microsoft.com/v1.0/me",
                headers={"Authorization": f"Bearer {token_acesso}"})
            with urllib.request.urlopen(req_perfil) as resp_perfil:
                dados_perfil = json.loads(resp_perfil.read().decode('utf-8'))
                email_usuario = dados_perfil.get('userPrincipalName', 'Conta Conectada')

            database.salvar_sessao(token_acesso, email_usuario)
            return render_template_string(_CALLBACK_TEMPLATE, titulo="Autenticacao Concluida com Sucesso!",
                mensagem="Credenciais registradas com seguranca.", instrucao=f"Conta: {email_usuario}")
    except urllib.error.HTTPError as e:
        log.error("OAuth HTTPError: %s %s", e.code, e.read().decode('utf-8', errors='replace'))
        return render_template_string(_CALLBACK_TEMPLATE, titulo="Erro na autenticacao",
            mensagem="Nao foi possivel obter token.", instrucao="Tente novamente."), 400
    except Exception as e:
        log.exception("OAuth erro: %s", e)
        return render_template_string(_CALLBACK_TEMPLATE, titulo="Erro na autenticacao",
            mensagem="Erro ao processar credenciais.", instrucao="Tente novamente."), 500


@app.route('/api/auth/status')
def auth_status():
    sessao = database.obter_sessao()
    if sessao:
        return jsonify({"logado": True, "conta": sessao["conta"]})
    return jsonify({"logado": False})


@app.route('/api/auth/logout', methods=['POST'])
def auth_logout():
    database.excluir_sessao()
    return jsonify({"mensagem": "Sessao encerrada."})


# -----------------------------------------------------------------------------
# Dashboard
# -----------------------------------------------------------------------------
@app.route('/api/dashboard')
def dashboard_metricas():
    return jsonify(database.obter_metricas_dashboard())


@app.route('/api/dashboard/dados')
def dashboard_dados_graficos():
    """Retorna dados para os gráficos do dashboard."""
    try:
        por_tipo = database.obter_dados_grafico_por_tipo()
        cidades = database.obter_dados_grafico_cidades(10)
        timeline = database.obter_dados_grafico_timeline(30)
        ultimos = database.obter_ultimos_clientes(5)
        return jsonify({
            "por_tipo": por_tipo,
            "cidades": {"labels": [c[0] for c in cidades], "valores": [c[1] for c in cidades]},
            "timeline": {"labels": [t[0] for t in timeline], "valores": [t[1] for t in timeline]},
            "ultimos_clientes": ultimos,
        })
    except Exception as e:
        log.exception("Erro ao obter dados dos graficos")
        return jsonify({"erro": str(e)}), 500


# -----------------------------------------------------------------------------
# Relatorios
# -----------------------------------------------------------------------------
@app.route('/api/relatorios')
def relatorios_filtrados():
    """Retorna clientes filtrados para relatórios."""
    tipo = request.args.get('tipo', '')
    cidade = request.args.get('cidade', '')
    data_inicio = request.args.get('data_inicio', '')
    data_fim = request.args.get('data_fim', '')
    clientes = database.listar_clientes_filtrados(tipo, cidade, data_inicio, data_fim)
    regioes = database.obter_perfis_regiao()
    return jsonify({"clientes": clientes, "total": len(clientes), "regioes": regioes})


@app.route('/api/relatorios/exportar')
def relatorios_exportar():
    """Exporta relatório filtrado como CSV ou XLSX."""
    tipo = request.args.get('tipo', '')
    cidade = request.args.get('cidade', '')
    data_inicio = request.args.get('data_inicio', '')
    data_fim = request.args.get('data_fim', '')
    formato = request.args.get('formato', 'csv')
    clientes = database.listar_clientes_filtrados(tipo, cidade, data_inicio, data_fim)

    cabecalhos = ["Nome", "CEP", "Rua", "Bairro", "Cidade", "Entrega/Retirada", "Observacoes", "Criado em"]
    campos = ['nome', 'cep', 'rua', 'bairro', 'cidade', 'entrega_retirada', 'observacoes', 'criado_em']

    if formato == 'xlsx':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Relatorio"
        fonte_cabecalho = Font(bold=True, color='FFFFFF', size=11)
        preenchimento_cabecalho = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
        alinhamento_central = Alignment(horizontal='center', vertical='center')
        borda_fina = Border(left=Side(style='thin', color='C0C0C0'), right=Side(style='thin', color='C0C0C0'),
            top=Side(style='thin', color='C0C0C0'), bottom=Side(style='thin', color='C0C0C0'))
        for col, titulo in enumerate(cabecalhos, start=1):
            cell = ws.cell(row=1, column=col, value=titulo)
            cell.font = fonte_cabecalho; cell.fill = preenchimento_cabecalho; cell.alignment = alinhamento_central; cell.border = borda_fina
        for i, c in enumerate(clientes, start=2):
            for col, campo in enumerate(campos, start=1):
                cell = ws.cell(row=i, column=col, value=c.get(campo, ''))
                cell.border = borda_fina; cell.font = Font(size=10); cell.alignment = alinhamento_central
        for col in range(1, len(cabecalhos) + 1):
            max_len = len(str(cabecalhos[col - 1]))
            for row in range(2, len(clientes) + 2):
                val = ws.cell(row=row, column=col).value
                if val:
                    max_len = max(max_len, min(len(str(val)), 40))
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max_len + 4
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        nome = f"relatorio_adri_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(buffer, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True, download_name=nome)
    else:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(cabecalhos)
        for c in clientes:
            writer.writerow([c.get(campo, '') for campo in campos])
        buffer.seek(0)
        nome = f"relatorio_adri_{time.strftime('%Y%m%d_%H%M%S')}.csv"
        return send_file(io.BytesIO(buffer.getvalue().encode('utf-8-sig')),
            mimetype='text/csv; charset=utf-8', as_attachment=True, download_name=nome)


# -----------------------------------------------------------------------------
# Backup
# -----------------------------------------------------------------------------
@app.route('/api/clientes/mapa')
def clientes_mapa():
    """Retorna clientes individuais com CEP para geolocalizacao no mapa."""
    return jsonify(database.obter_clientes_para_mapa())


@app.route('/api/backup/json')
def backup_json():
    clientes = database.listar_clientes()
    dados = {"versao": "1.0", "data_exportacao": time.strftime('%Y-%m-%dT%H:%M:%S'), "clientes": clientes}
    buffer = io.BytesIO()
    buffer.write(json.dumps(dados, ensure_ascii=False, indent=2).encode('utf-8'))
    buffer.seek(0)
    return send_file(buffer, mimetype='application/json', as_attachment=True,
        download_name=f"backup_adri_poltronas_{time.strftime('%Y%m%d_%H%M%S')}.json")


@app.route('/api/backup/sqlite')
def backup_sqlite():
    caminho = database.obter_caminho_banco_absoluto()
    if not os.path.exists(caminho):
        return jsonify({"erro": "Banco de dados nao encontrado."}), 404
    return send_file(caminho, mimetype='application/x-sqlite3', as_attachment=True,
        download_name=f"backup_adri_poltronas_{time.strftime('%Y%m%d_%H%M%S')}.db")


# -----------------------------------------------------------------------------
# Importacao
# -----------------------------------------------------------------------------
@app.route('/api/clientes/importar', methods=['POST'])
def importar_clientes():
    if 'arquivo' not in request.files:
        return jsonify({"erro": "Nenhum arquivo enviado."}), 400
    arquivo = request.files['arquivo']
    nome = arquivo.filename.lower()
    clientes = []

    try:
        if nome.endswith('.csv'):
            conteudo = arquivo.read().decode('utf-8-sig')
            leitor = csv.DictReader(io.StringIO(conteudo), delimiter=',')
            for linha in leitor:
                clientes.append({
                    'nome': linha.get('nome') or linha.get('Nome') or linha.get('name') or '',
                    'cep': linha.get('cep') or linha.get('CEP') or linha.get('zip') or '',
                    'rua': linha.get('rua') or linha.get('Rua') or linha.get('address') or '',
                    'bairro': linha.get('bairro') or linha.get('Bairro') or linha.get('neighborhood') or '',
                    'cidade': linha.get('cidade') or linha.get('Cidade') or linha.get('city') or '',
                    'entrega_retirada': linha.get('entrega_retirada') or linha.get('Entrega/Retirada') or 'Entrega',
                    'observacoes': linha.get('observacoes') or linha.get('Observações') or '',
                })
        elif nome.endswith('.json'):
            dados = json.loads(arquivo.read().decode('utf-8'))
            clientes = dados if isinstance(dados, list) else dados.get('clientes', [])
        elif nome.endswith('.xlsx') or nome.endswith('.xls'):
            wb = openpyxl.load_workbook(arquivo)
            ws = wb.active
            headers = [str(cell.value or '').strip().lower() for cell in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                d = {headers[i]: str(row[i] or '') for i in range(len(headers)) if i < len(row)}
                clientes.append({
                    'nome': d.get('nome', ''),
                    'cep': d.get('cep', d.get('zip', '')),
                    'rua': d.get('rua', d.get('logradouro', '')),
                    'bairro': d.get('bairro', d.get('neighborhood', '')),
                    'cidade': d.get('cidade', d.get('city', '')),
                    'entrega_retirada': d.get('entrega_retirada', d.get('tipo', 'Entrega')),
                    'observacoes': d.get('observacoes', ''),
                })
        else:
            return jsonify({"erro": "Formato nao suportado. Use CSV, JSON ou XLSX."}), 400
    except Exception as e:
        log.exception("Erro ao ler arquivo: %s", e)
        return jsonify({"erro": f"Erro ao ler arquivo: {str(e)}"}), 400

    if not clientes:
        return jsonify({"erro": "Nenhum cliente encontrado no arquivo."}), 400

    clientes_validados = []
    erros_validacao = []
    for idx, c in enumerate(clientes):
        dados_validados, erros = validators.validar_cliente(c)
        if erros:
            erros_validacao.append(f"{c.get('nome', f'linha {idx+1}')}: {'; '.join(erros)}")
        else:
            clientes_validados.append(dados_validados)

    if not clientes_validados:
        return jsonify({"erro": "Nenhum cliente valido para importar.", "detalhes": erros_validacao}), 400

    resultado = database.importar_clientes_em_lote(clientes_validados)
    resposta = {"mensagem": "Importacao concluida!", "sucessos": resultado["sucessos"],
        "duplicados": resultado["duplicados"], "erros": resultado["erros"], "total": len(clientes)}
    if erros_validacao:
        resposta["detalhes_validacao"] = erros_validacao
    return jsonify(resposta)


# -----------------------------------------------------------------------------
# Exportacao Excel (OneDrive)
# -----------------------------------------------------------------------------
@app.route('/api/excel/enviar', methods=['POST'])
def enviar_excel_onedrive():
    sessao = database.obter_sessao()
    if not sessao:
        return jsonify({"erro": "Usuario nao autenticado."}), 401
    token = sessao["token"]
    dados_requisicao = request.get_json(silent=True) or {}
    ids_selecionados = dados_requisicao.get('ids', [])
    nome_personalizado = dados_requisicao.get('nome_arquivo', '').strip()
    if not ids_selecionados:
        return jsonify({"erro": "Nenhum cliente selecionado."}), 400
    todos_clientes = database.listar_clientes()
    clientes = [c for c in todos_clientes if c['id'] in ids_selecionados]
    if not clientes:
        return jsonify({"erro": "Clientes nao encontrados."}), 404
    if not nome_personalizado:
        nome_personalizado = "Clientes_AdriPoltronas"
    nome_base = nome_personalizado.rsplit('.', 1)[0] if '.' in nome_personalizado else nome_personalizado
    nome_arquivo = f"{nome_base}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clientes"
    fonte_cabecalho = Font(bold=True, color='FFFFFF', size=11)
    preenchimento_cabecalho = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    alinhamento_central = Alignment(horizontal='center', vertical='center')
    borda_fina = Border(left=Side(style='thin', color='C0C0C0'), right=Side(style='thin', color='C0C0C0'),
        top=Side(style='thin', color='C0C0C0'), bottom=Side(style='thin', color='C0C0C0'))
    cabecalhos = ["Nome", "CEP", "Rua", "Bairro", "Cidade", "Entrega/Retirada", "Observações"]
    for col, titulo in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.font = fonte_cabecalho; cell.fill = preenchimento_cabecalho; cell.alignment = alinhamento_central; cell.border = borda_fina

    for i, c in enumerate(clientes, start=2):
        for col, campo in enumerate(['nome', 'cep', 'rua', 'bairro', 'cidade', 'entrega_retirada', 'observacoes'], start=1):
            cell = ws.cell(row=i, column=col, value=c.get(campo, ''))
            cell.border = borda_fina
            cell.font = Font(size=10)
            cell.alignment = alinhamento_central

    for col in range(1, len(cabecalhos) + 1):
        max_len = len(str(cabecalhos[col - 1]))
        for row in range(2, len(clientes) + 2):
            val = ws.cell(row=row, column=col).value
            if val:
                max_len = max(max_len, min(len(str(val)), 40))
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max_len + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    corpo_binario = buffer.read()

    url_upload = f"https://graph.microsoft.com/v1.0/me/drive/root:/{urllib.parse.quote(nome_arquivo)}:/content"
    req_upload = urllib.request.Request(url_upload, data=corpo_binario,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        method='PUT')
    try:
        with urllib.request.urlopen(req_upload) as resp:
            retorno = json.loads(resp.read().decode('utf-8'))
            web_url = retorno.get('webUrl')
        if web_url:
            webbrowser.open(web_url)
        return jsonify({"mensagem": "Planilha enviada!", "url": web_url}), 200
    except urllib.error.HTTPError as e:
        erro_corpo = e.read().decode('utf-8', errors='replace')
        log.error("Graph HTTP %s: %s", e.code, erro_corpo)
        if e.code == 401:
            database.excluir_sessao()
            return jsonify({"erro": "Sessao expirada. Reconecte."}), 401
        return jsonify({"erro": f"Erro {e.code} ao enviar para Excel."}), 500
    except urllib.error.URLError as e:
        log.error("Conexao com Graph: %s", e.reason)
        return jsonify({"erro": "Falha de conexao com a Microsoft."}), 502


# -----------------------------------------------------------------------------
# IA - Sugestao de Bairro por CEP
# -----------------------------------------------------------------------------
@app.route('/api/ia/sugestao-bairro', methods=['POST'])
def ia_sugestao_bairro():
    dados = request.get_json(silent=True) or {}
    cep = (dados.get('cep') or '').replace('-', '').strip()
    if len(cep) != 8 or not cep.isdigit():
        return jsonify({"erro": "CEP invalido."}), 400
    try:
        url_viacep = f"https://viacep.com.br/ws/{cep}/json/"
        with urllib.request.urlopen(urllib.request.Request(url_viacep), timeout=5) as resp:
            dados_cep = json.loads(resp.read().decode('utf-8'))
            if not dados_cep.get('erro'):
                return jsonify({"bairro": dados_cep.get('bairro', ''), "cidade": dados_cep.get('localidade', ''),
                    "rua": dados_cep.get('logradouro', ''), "uf": dados_cep.get('uf', ''), "fonte": "ViaCEP"})
    except Exception as e:
        log.warning("ViaCEP falhou: %s", e)
    if not GEMINI_API_KEY:
        return jsonify({"erro": "Nao foi possivel buscar dados para este CEP."}), 500
    try:
        prompt = f"Para o CEP brasileiro {cep[:5]}-{cep[5:]}, quais sao os bairros mais provaveis? Responda APENAS com JSON: {{\"bairro\": \"nome do bairro\"}}"
        payload = json.dumps({"contents": [{"parts": [{"text": prompt}]}], "generationConfig": {"temperature": 0.3, "maxOutputTokens": 200}}).encode()
        req = urllib.request.Request(f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GEMINI_API_KEY}",
            data=payload, headers={"Content-Type": "application/json"}, method='POST')
        with urllib.request.urlopen(req, timeout=10) as resp:
            resultado = json.loads(resp.read().decode('utf-8'))
            texto = resultado.get('candidates', [{}])[0].get('content', {}).get('parts', [{}])[0].get('text', '')
            match = re.search(r'\{[^}]+\}', texto)
            if match:
                return jsonify({"bairro": json.loads(match.group()).get('bairro', ''), "fonte": "Gemini AI"})
    except Exception as e:
        log.warning("Gemini falhou: %s", e)
    return jsonify({"erro": "Nao foi possivel buscar dados."}), 500


# -----------------------------------------------------------------------------
# Pagamentos
# -----------------------------------------------------------------------------
@app.route('/api/pagamentos', methods=['GET'])
def listar_pagamentos():
    cliente_id = request.args.get('cliente_id')
    status = request.args.get('status')
    pagamentos = database.listar_pagamentos(cliente_id=int(cliente_id) if cliente_id else None, status=status)
    return jsonify(pagamentos)


@app.route('/api/pagamentos/metricas')
def pagamentos_metricas():
    return jsonify(database.obter_metricas_pagamentos())


@app.route('/api/pagamentos/asaas-status')
def asaas_status():
    return jsonify({
        "configurado": asaas.esta_configurado(),
        "ambiente": asaas.ambiente(),
    })


@app.route('/api/pagamentos', methods=['POST'])
def criar_pagamento():
    pagamento, erros = validators.validar_pagamento(request.get_json(silent=True) or {})
    if erros:
        return jsonify({"erro": "Dados invalidos.", "detalhes": erros}), 400
    pagamento_id = database.salvar_pagamento(
        pagamento['cliente_id'], pagamento['cliente_nome'], pagamento['valor'],
        pagamento['descricao'], pagamento['data_vencimento'], pagamento['tipo_pagamento'],
    )
    return jsonify({"mensagem": "Pagamento registrado!", "id": pagamento_id}), 201


@app.route('/api/pagamentos/<int:pagamento_id>', methods=['PUT'])
def atualizar_pagamento(pagamento_id):
    dados = request.get_json(silent=True) or {}
    pagamento, erros = validators.validar_pagamento(dados)
    if erros:
        return jsonify({"erro": "Dados invalidos.", "detalhes": erros}), 400
    existente = database.obter_pagamento_por_id(pagamento_id)
    if not existente:
        return jsonify({"erro": "Pagamento nao encontrado."}), 404
    linhas = database.atualizar_pagamento(pagamento_id, **pagamento)
    return jsonify({"mensagem": "Atualizado!", "linhas": linhas}), 200


@app.route('/api/pagamentos/<int:pagamento_id>', methods=['DELETE'])
def excluir_pagamento(pagamento_id):
    existente = database.obter_pagamento_por_id(pagamento_id)
    if not existente:
        return jsonify({"erro": "Pagamento nao encontrado."}), 404
    if existente.get('asaas_payment_id') and asaas.esta_configurado():
        resultado = asaas.cancelar_pagamento(existente['asaas_payment_id'])
        if resultado.get('_erro'):
            log.warning("Falha ao cancelar no Asaas: %s", resultado)
    linhas = database.excluir_pagamento(pagamento_id)
    if linhas == 0:
        return jsonify({"erro": "Nao foi possivel excluir (so e possivel cancelar pagamentos pendentes)."}), 400
    return jsonify({"mensagem": "Pagamento cancelado!"}), 200


@app.route('/api/pagamentos/<int:pagamento_id>/confirmar', methods=['POST'])
def confirmar_pagamento(pagamento_id):
    existente = database.obter_pagamento_por_id(pagamento_id)
    if not existente:
        return jsonify({"erro": "Pagamento nao encontrado."}), 404
    database.atualizar_pagamento(pagamento_id, status='Recebido',
        data_pagamento=time.strftime('%Y-%m-%d'))
    return jsonify({"mensagem": "Pagamento confirmado!"}), 200


@app.route('/api/pagamentos/gerar-cobranca', methods=['POST'])
def gerar_cobranca():
    """Gera cobranca Pix ou Boleto no Asaas e retorna os dados."""
    if not asaas.esta_configurado():
        return jsonify({"erro": "Integracao Asaas nao configurada. Defina ASAAS_API_KEY."}), 503
    dados = request.get_json(silent=True) or {}
    pagamento_id = dados.get('pagamento_id')
    if not pagamento_id:
        return jsonify({"erro": "pagamento_id e obrigatorio."}), 400
    pagamento = database.obter_pagamento_por_id(int(pagamento_id))
    if not pagamento:
        return jsonify({"erro": "Pagamento nao encontrado."}), 404
    cliente = database.obter_cliente_por_id(pagamento['cliente_id'])
    if not cliente:
        return jsonify({"erro": "Cliente nao encontrado."}), 404
    asaas_customer_id = database.obter_ou_criar_cliente_asaas(
        pagamento['cliente_id'], cliente['nome'],
        cidade=cliente.get('cidade'), bairro=cliente.get('bairro'),
        rua=cliente.get('rua'), cep=cliente.get('cep'),
    )
    if not asaas_customer_id:
        return jsonify({"erro": "Falha ao criar cliente no Asaas."}), 502
    tipo = pagamento.get('tipo_pagamento', 'Pix')
    if tipo == 'Pix':
        resultado = asaas.criar_cobranca_pix(
            asaas_customer_id, pagamento['valor'], pagamento['data_vencimento'],
            pagamento.get('descricao', ''),
        )
    elif tipo == 'Boleto':
        resultado = asaas.criar_cobranca_boleto(
            asaas_customer_id, pagamento['valor'], pagamento['data_vencimento'],
            pagamento.get('descricao', ''),
        )
    else:
        return jsonify({"erro": "Tipo de pagamento invalido."}), 400
    if resultado.get('_erro'):
        return jsonify({"erro": "Falha ao gerar cobranca no Asaas.", "detalhes": resultado}), 502
    asaas_payment_id = resultado.get('id', '')
    update_fields = {'asaas_payment_id': asaas_payment_id}
    if tipo == 'Boleto':
        update_fields['boleto_url'] = resultado.get('bankSlipUrl', '')
        update_fields['nosso_numero'] = resultado.get('nossoNumero', '')
    database.atualizar_pagamento(pagamento_id, **update_fields)
    resposta = {
        "mensagem": f"Cobranca {tipo} gerada com sucesso!",
        "asaas_payment_id": asaas_payment_id,
        "tipo": tipo,
    }
    if tipo == 'Pix':
        pix_data = asaas.obter_qr_code_pix(asaas_payment_id)
        if not pix_data.get('_erro'):
            database.atualizar_pagamento(pagamento_id,
                pix_payload=pix_data.get('payload', ''),
                pix_qr_code=pix_data.get('encodedImage', ''),
            )
            resposta['pix_payload'] = pix_data.get('payload', '')
            resposta['pix_qr_code'] = pix_data.get('encodedImage', '')
        else:
            log.warning("Falha ao obter QR Code Pix: %s", pix_data)
    elif tipo == 'Boleto':
        resposta['boleto_url'] = resultado.get('bankSlipUrl', '')
        resposta['nosso_numero'] = resultado.get('nossoNumero', '')
    return jsonify(resposta), 201


@app.route('/api/pagamentos/<int:pagamento_id>/cobranca-dados')
def obter_dados_cobranca(pagamento_id):
    """Retorna dados da cobranca (QR Code Pix ou URL Boleto)."""
    pagamento = database.obter_pagamento_por_id(pagamento_id)
    if not pagamento:
        return jsonify({"erro": "Pagamento nao encontrado."}), 404
    resposta = {
        "id": pagamento['id'],
        "tipo_pagamento": pagamento.get('tipo_pagamento', 'Pix'),
        "valor": pagamento['valor'],
        "status": pagamento['status'],
        "asaas_payment_id": pagamento.get('asaas_payment_id', ''),
    }
    if pagamento.get('tipo_pagamento') == 'Pix':
        resposta['pix_payload'] = pagamento.get('pix_payload', '')
        resposta['pix_qr_code'] = pagamento.get('pix_qr_code', '')
    elif pagamento.get('tipo_pagamento') == 'Boleto':
        resposta['boleto_url'] = pagamento.get('boleto_url', '')
        resposta['nosso_numero'] = pagamento.get('nosso_numero', '')
    return jsonify(resposta)


@app.route('/api/pagamentos/webhook', methods=['POST'])
def webhook_asaas():
    """Endpoint para webhooks do Asaas (confirmacao de pagamento)."""
    dados = request.get_json(silent=True) or {}
    log.info("[WEBHOOK] Recebido: %s", json.dumps(dados, ensure_ascii=False)[:500])
    tipo_evento = dados.get('event', '')
    pagamento_asaas = dados.get('payment', {})
    asaas_payment_id = pagamento_asaas.get('id', '')
    if not asaas_payment_id:
        return jsonify({"mensagem": "OK"}), 200
    if tipo_evento in ('PAYMENT_RECEIVED', 'PAYMENT_CONFIRMED'):
        with database.conectar() as conexao:
            pagamento = conexao.execute(
                'SELECT * FROM pagamentos WHERE asaas_payment_id=?', (asaas_payment_id,)
            ).fetchone()
            if pagamento:
                conexao.execute(
                    'UPDATE pagamentos SET status="Recebido", data_pagamento=datetime("now", "localtime") WHERE asaas_payment_id=?',
                    (asaas_payment_id,)
                )
                conexao.execute(
                    'INSERT INTO historico (cliente_id, cliente_nome, acao, detalhes) VALUES (?, ?, ?, ?)',
                    (pagamento['cliente_id'], pagamento['cliente_nome'], 'PAGAMENTO',
                     f'Pagamento confirmado via webhook: R$ {pagamento["valor"]:.2f}'),
                )
                log.info("[WEBHOOK] Pagamento %s confirmado via Asaas.", asaas_payment_id)
    elif tipo_evento == 'PAYMENT_OVERDUE':
        with database.conectar() as conexao:
            conexao.execute(
                'UPDATE pagamentos SET status="Atrasado" WHERE asaas_payment_id=? AND status="Pendente"',
                (asaas_payment_id,)
            )
    elif tipo_evento == 'PAYMENT_DELETED':
        with database.conectar() as conexao:
            conexao.execute(
                'UPDATE pagamentos SET status="Cancelado" WHERE asaas_payment_id=?',
                (asaas_payment_id,)
            )
    return jsonify({"mensagem": "OK"}), 200


# -----------------------------------------------------------------------------
# Notificacoes Inteligentes
# -----------------------------------------------------------------------------
@app.route('/api/notificacoes')
def api_obter_notificacoes():
    """Retorna notificacoes inteligentes do sistema."""
    try:
        notifs = database.obter_notificacoes()
        return jsonify({"notificacoes": notifs, "total": len(notifs)})
    except Exception as e:
        log.exception("Erro ao obter notificacoes")
        return jsonify({"notificacoes": [], "total": 0, "erro": str(e)}), 500


@app.route('/api/entregas', methods=['GET'])
def listar_entregas():
    """Lista entregas agendadas."""
    status = request.args.get('status')
    return jsonify(database.listar_entregas_agendadas(status=status))


@app.route('/api/entregas', methods=['POST'])
def criar_entrega():
    """Agenda uma nova entrega."""
    dados = request.get_json(silent=True) or {}
    cliente_id = dados.get('cliente_id')
    cliente_nome = dados.get('cliente_nome', '')
    data_entrega = dados.get('data_entrega', '')
    hora_entrega = dados.get('hora_entrega', '')
    descricao = dados.get('descricao', '')
    if not cliente_id or not data_entrega:
        return jsonify({"erro": "cliente_id e data_entrega sao obrigatorios."}), 400
    entrega_id = database.agendar_entrega(cliente_id, cliente_nome, data_entrega, hora_entrega, descricao)
    return jsonify({"mensagem": "Entrega agendada!", "id": entrega_id}), 201


@app.route('/api/entregas/<int:entrega_id>/status', methods=['PUT'])
def atualizar_status_entrega(entrega_id):
    """Atualiza status de uma entrega."""
    dados = request.get_json(silent=True) or {}
    novo_status = dados.get('status', '')
    if novo_status not in ('Pendente', 'Em andamento', 'Concluida', 'Cancelada'):
        return jsonify({"erro": "Status invalido."}), 400
    linhas = database.atualizar_status_entrega(entrega_id, novo_status)
    if linhas == 0:
        return jsonify({"erro": "Entrega nao encontrada."}), 404
    return jsonify({"mensagem": "Status atualizado!"}), 200


# -----------------------------------------------------------------------------
# CRUD Clientes
# -----------------------------------------------------------------------------
@app.route('/api/clientes', methods=['GET'])
def obtener_clientes():
    busca = request.args.get('busca', '')
    return jsonify(database.listar_clientes(busca))


@app.route('/api/clientes', methods=['POST'])
def adicionar_cliente():
    cliente, erros = validators.validar_cliente(request.get_json(silent=True) or {})
    if erros:
        return jsonify({"erro": "Dados invalidos.", "detalhes": erros}), 400
    database.salvar_cliente(cliente['nome'], cliente['cep'], cliente['rua'], cliente['bairro'],
        cliente['cidade'], cliente['entrega_retirada'], cliente.get('observacoes', ''))
    return jsonify({"mensagem": "Salvo"}), 201


@app.route('/api/clientes/<int:id_cliente>', methods=['PUT'])
def atualizar_cliente(id_cliente):
    cliente, erros = validators.validar_cliente(request.get_json(silent=True) or {})
    if erros:
        return jsonify({"erro": "Dados invalidos.", "detalhes": erros}), 400
    database.atualizar_cliente(id_cliente, cliente['nome'], cliente['cep'], cliente['rua'],
        cliente['bairro'], cliente['cidade'], cliente['entrega_retirada'], cliente.get('observacoes', ''))
    return jsonify({"mensagem": "Atualizado"}), 200


@app.route('/api/clientes/<int:id_cliente>', methods=['DELETE'])
def excluir_cliente(id_cliente):
    linhas = database.excluir_cliente(id_cliente)
    if linhas == 0:
        return jsonify({"erro": "Cliente nao encontrado."}), 404
    return jsonify({"mensagem": "Excluido"}), 200


if __name__ == '__main__':
    database.criar_tabela()
    app.run(host='localhost', port=5000)
