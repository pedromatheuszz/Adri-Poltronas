# backend/blueprints/excel.py
"""Blueprint de exportação Excel para OneDrive."""
import io
import json
import urllib.parse
import urllib.request
import webbrowser
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import Blueprint, request, jsonify

import database
import config

log = logging.getLogger("adri.excel")

excel_bp = Blueprint('excel', __name__, url_prefix='/api/excel')


@excel_bp.route('/enviar', methods=['POST'])
def enviar_excel_onedrive():
    sessao = database.obter_sessao()
    if not sessao:
        return jsonify({"erro": "Usuario nao autenticado."}), 401
    token = sessao["token"]
    dados_requisicao = request.get_json(silent=True) or {}
    ids_selecionados = dados_requisicao.get('ids', [])
    nome_personalizado = dados_requisicao.get('nome_arquivo', '').strip()
    if not ids_selecionados:
        return jsonify({"erro": "Nenhum cliente selecionado."}), 400
    todos_clientes = database.listar_clientes()
    clientes = [c for c in todos_clientes if c['id'] in ids_selecionados]
    if not clientes:
        return jsonify({"erro": "Clientes nao encontrados."}), 404
    if not nome_personalizado:
        nome_personalizado = "Clientes_AdriPoltronas"
    nome_base = nome_personalizado.rsplit('.', 1)[0] if '.' in nome_personalizado else nome_personalizado
    nome_arquivo = f"{nome_base}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clientes"
    fonte_cabecalho = Font(bold=True, color='FFFFFF', size=11)
    preenchimento_cabecalho = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    alinhamento_central = Alignment(horizontal='center', vertical='center')
    borda_fina = Border(left=Side(style='thin', color='C0C0C0'), right=Side(style='thin', color='C0C0C0'),
        top=Side(style='thin', color='C0C0C0'), bottom=Side(style='thin', color='C0C0C0'))
    cabecalhos = ["Nome", "CEP", "Rua", "Bairro", "Cidade", "Entrega/Retirada", "Observações"]
    for col, titulo in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.font = fonte_cabecalho; cell.fill = preenchimento_cabecalho; cell.alignment = alinhamento_central; cell.border = borda_fina

    for i, c in enumerate(clientes, start=2):
        for col, campo in enumerate(['nome', 'cep', 'rua', 'bairro', 'cidade', 'entrega_retirada', 'observacoes'], start=1):
            cell = ws.cell(row=i, column=col, value=c.get(campo, ''))
            cell.border = borda_fina
            cell.font = Font(size=10)
            cell.alignment = alinhamento_central

    for col in range(1, len(cabecalhos) + 1):
        max_len = len(str(cabecalhos[col - 1]))
        for row in range(2, len(clientes) + 2):
            val = ws.cell(row=row, column=col).value
            if val:
                max_len = max(max_len, min(len(str(val)), 40))
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max_len + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    corpo_binario = buffer.read()

    url_upload = f"https://graph.microsoft.com/v1.0/me/drive/root:/{urllib.parse.quote(nome_arquivo)}:/content"
    req_upload = urllib.request.Request(url_upload, data=corpo_binario,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        method='PUT')
    try:
        with urllib.request.urlopen(req_upload) as resp:
            retorno = json.loads(resp.read().decode('utf-8'))
            web_url = retorno.get('webUrl')
        if web_url:
            webbrowser.open(web_url)
        return jsonify({"mensagem": "Planilha enviada!", "url": web_url}), 200
    except urllib.error.HTTPError as e:
        erro_corpo = e.read().decode('utf-8', errors='replace')
        log.error("Graph HTTP %s: %s", e.code, erro_corpo)
        if e.code == 401:
            database.excluir_sessao()
            return jsonify({"erro": "Sessao expirada. Reconecte."}), 401
        return jsonify({"erro": f"Erro {e.code} ao enviar para Excel."}), 500
    except urllib.error.URLError as e:
        log.error("Conexao com Graph: %s", e.reason)
        return jsonify({"erro": "Falha de conexao com a Microsoft."}), 502