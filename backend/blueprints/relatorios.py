# backend/blueprints/relatorios.py
"""Blueprint de relatórios."""
import time
import io
import csv
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import Blueprint, request, jsonify, send_file

import database

log = logging.getLogger("adri.relatorios")

relatorios_bp = Blueprint('relatorios', __name__, url_prefix='/api/relatorios')


@relatorios_bp.route('', methods=['GET'])
def relatorios_filtrados():
    """Retorna clientes filtrados para relatórios."""
    tipo = request.args.get('tipo', '')
    cidade = request.args.get('cidade', '')
    data_inicio = request.args.get('data_inicio', '')
    data_fim = request.args.get('data_fim', '')
    clientes = database.listar_clientes_filtrados(tipo, cidade, data_inicio, data_fim)
    regioes = database.obter_perfis_regiao()
    return jsonify({"clientes": clientes, "total": len(clientes), "regioes": regioes})


@relatorios_bp.route('/exportar')
def relatorios_exportar():
    """Exporta relatório filtrado como CSV ou XLSX."""
    tipo = request.args.get('tipo', '')
    cidade = request.args.get('cidade', '')
    data_inicio = request.args.get('data_inicio', '')
    data_fim = request.args.get('data_fim', '')
    formato = request.args.get('formato', 'csv')
    clientes = database.listar_clientes_filtrados(tipo, cidade, data_inicio, data_fim)

    cabecalhos = ["Nome", "CEP", "Rua", "Bairro", "Cidade", "Entrega/Retirada", "Observacoes", "Criado em"]
    campos = ['nome', 'cep', 'rua', 'bairro', 'cidade', 'entrega_retirada', 'observacoes', 'criado_em']

    if formato == 'xlsx':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Relatorio"
        fonte_cabecalho = Font(bold=True, color='FFFFFF', size=11)
        preenchimento_cabecalho = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
        alinhamento_central = Alignment(horizontal='center', vertical='center')
        borda_fina = Border(left=Side(style='thin', color='C0C0C0'), right=Side(style='thin', color='C0C0C0'),
            top=Side(style='thin', color='C0C0C0'), bottom=Side(style='thin', color='C0C0C0'))
        for col, titulo in enumerate(cabecalhos, start=1):
            cell = ws.cell(row=1, column=col, value=titulo)
            cell.font = fonte_cabecalho; cell.fill = preenchimento_cabecalho; cell.alignment = alinhamento_central; cell.border = borda_fina
        for i, c in enumerate(clientes, start=2):
            for col, campo in enumerate(campos, start=1):
                cell = ws.cell(row=i, column=col, value=c.get(campo, ''))
                cell.border = borda_fina; cell.font = Font(size=10); cell.alignment = alinhamento_central
        for col in range(1, len(cabecalhos) + 1):
            max_len = len(str(cabecalhos[col - 1]))
            for row in range(2, len(clientes) + 2):
                val = ws.cell(row=row, column=col).value
                if val:
                    max_len = max(max_len, min(len(str(val)), 40))
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max_len + 4
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        nome = f"relatorio_adri_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(buffer, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True, download_name=nome)
    else:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(cabecalhos)
        for c in clientes:
            writer.writerow([c.get(campo, '') for campo in campos])
        buffer.seek(0)
        nome = f"relatorio_adri_{time.strftime('%Y%m%d_%H%M%S')}.csv"
        return send_file(io.BytesIO(buffer.getvalue().encode('utf-8-sig')),
            mimetype='text/csv; charset=utf-8', as_attachment=True, download_name=nome)