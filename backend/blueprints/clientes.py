# backend/blueprints/clientes.py
"""Blueprint de clientes (CRUD, importação, backup, mapa)."""
import time
import io
import csv
import json
import logging
from functools import wraps

import openpyxl
from flask import Blueprint, request, jsonify, send_file

import database
import security
import validators
from erp.core import idempotencia

log = logging.getLogger("adri.clientes")

# Teto de registros por importacao.
#
# MAX_CONTENT_LENGTH ja limita o corpo da requisicao, mas 10 MB de CSV ainda
# sao centenas de milhares de linhas -- cada uma virando um INSERT dentro de
# uma unica transacao SQLite, com o app travado enquanto isso.
MAX_LINHAS_IMPORTACAO = 5000

clientes_bp = Blueprint('clientes', __name__, url_prefix='/api/clientes')


def idempotente(rota):
    """
    Faz a rota lembrar do que ja processou, por chave de operacao.

    Existe por causa da fila offline. Ela reenvia sempre que a tentativa
    anterior nao confirmou -- e "nao confirmou" e ambiguo: pode ser que o
    servidor nunca tenha recebido, ou que tenha recebido, gravado, e a
    resposta ter se perdido no caminho de volta. Do lado do navegador as duas
    situacoes sao o mesmo fetch que rejeita, e reenviar um POST de criacao
    produzia dois clientes identicos.

    Sem cabecalho `Idempotency-Key` a rota se comporta como antes: o
    mecanismo e para quem envia da fila, nao para a tela.
    """
    @wraps(rota)
    def envelope(*args, **kwargs):
        chave = (request.headers.get(idempotencia.NOME_DO_HEADER) or '').strip()
        if not idempotencia.chave_valida(chave):
            return rota(*args, **kwargs)

        lembrada = idempotencia.resposta_registrada(database.conectar, chave)
        if lembrada is not None:
            status, corpo = lembrada
            resposta = jsonify(corpo)
            resposta.status_code = status
            # Deixa explicito, para quem depurar, que este 201 nao criou nada.
            resposta.headers['X-Idempotencia'] = 'repetida'
            return resposta

        resultado = rota(*args, **kwargs)
        corpo, status = _corpo_e_status(resultado)
        # So sucesso e lembrado: reenviar um 400 de validacao produz o mesmo
        # 400, sem efeito colateral, e guardar erro so faria a tabela crescer.
        if corpo is not None and 200 <= status < 300:
            idempotencia.registrar(
                database.conectar, chave, request.path, request.method, status, corpo
            )
        return resultado

    return envelope


def _corpo_e_status(resultado):
    """Extrai (dict, status) do que uma view devolveu."""
    resposta, status = resultado if isinstance(resultado, tuple) else (resultado, 200)
    try:
        return resposta.get_json(), int(status)
    except (AttributeError, TypeError, ValueError):
        return None, int(status) if isinstance(status, int) else 200



@clientes_bp.route('', methods=['GET'])
def listar_clientes():
    busca = request.args.get('busca', '')
    return jsonify(database.listar_clientes(busca))


@clientes_bp.route('', methods=['POST'])
@idempotente
def adicionar_cliente():
    cliente, erros = validators.validar_cliente(request.get_json(silent=True) or {})
    if erros:
        return jsonify({"erro": "Dados invalidos.", "detalhes": erros}), 400
    # O `id` na resposta nao e enfeite: a fila offline cria o cliente com um
    # id temporario negativo e precisa do id real para remapear as operacoes
    # que vierem depois na mesma fila. Sem ele, a edicao seguinte era enviada
    # contra o id negativo, tomava 404 e era descartada em silencio.
    cliente_id = database.salvar_cliente(
        cliente['nome'], cliente['cep'], cliente['rua'], cliente['bairro'],
        cliente['cidade'], cliente['entrega_retirada'], cliente.get('observacoes', '')
    )
    return jsonify({"mensagem": "Cliente cadastrado!", "id": cliente_id}), 201


@clientes_bp.route('/<int:id_cliente>', methods=['PUT'])
@idempotente
def atualizar_cliente(id_cliente):
    cliente, erros = validators.validar_cliente(request.get_json(silent=True) or {})
    if erros:
        return jsonify({"erro": "Dados invalidos.", "detalhes": erros}), 400
    linhas = database.atualizar_cliente(
        id_cliente, cliente['nome'], cliente['cep'], cliente['rua'],
        cliente['bairro'], cliente['cidade'], cliente['entrega_retirada'],
        cliente.get('observacoes', '')
    )
    if linhas == 0:
        return jsonify({"erro": "Cliente nao encontrado."}), 404
    return jsonify({"mensagem": "Cliente atualizado!"}), 200


@clientes_bp.route('/<int:id_cliente>', methods=['DELETE'])
@idempotente
def excluir_cliente(id_cliente):
    linhas = database.excluir_cliente(id_cliente)
    if linhas == 0:
        return jsonify({"erro": "Cliente nao encontrado."}), 404
    return jsonify({"mensagem": "Cliente excluido!"}), 200


@clientes_bp.route('/importar', methods=['POST'])
def importar_clientes():
    if 'arquivo' not in request.files:
        return jsonify({"erro": "Nenhum arquivo enviado."}), 400
    arquivo = request.files['arquivo']

    # Extensao, tamanho e arquivo vazio antes de qualquer parsing: o openpyxl
    # abre zip, e nao convem entregar a ele um arquivo nao verificado.
    valido, erro_upload = security.validar_upload(arquivo)
    if not valido:
        return jsonify({"erro": erro_upload}), 400

    nome = (arquivo.filename or '').lower()
    clientes = []

    try:
        if nome.endswith('.csv'):
            conteudo = arquivo.read().decode('utf-8-sig')
            leitor = csv.DictReader(io.StringIO(conteudo), delimiter=',')
            for linha in leitor:
                clientes.append({
                    'nome': linha.get('nome') or linha.get('Nome') or linha.get('name') or '',
                    'cep': linha.get('cep') or linha.get('CEP') or linha.get('zip') or '',
                    'rua': linha.get('rua') or linha.get('Rua') or linha.get('address') or '',
                    'bairro': linha.get('bairro') or linha.get('Bairro') or linha.get('neighborhood') or '',
                    'cidade': linha.get('cidade') or linha.get('Cidade') or linha.get('city') or '',
                    'entrega_retirada': linha.get('entrega_retirada') or linha.get('Entrega/Retirada') or 'Entrega',
                    'observacoes': (linha.get('observacoes') or linha.get('Observacoes')
                                    or linha.get('Observações') or ''),
                })
        elif nome.endswith('.json'):
            dados = json.loads(arquivo.read().decode('utf-8'))
            clientes = dados if isinstance(dados, list) else dados.get('clientes', [])
        elif nome.endswith('.xlsx'):
            # `.xls` (Excel antigo, formato BIFF) nao chega ate aqui: o
            # openpyxl so le OOXML, e `security.validar_upload` recusa a
            # extensao com a instrucao de converter. Antes ela era aceita, o
            # arquivo ia direto para `load_workbook` e o usuario recebia
            # "nao foi possivel ler o arquivo" sem saber o que fazer.
            wb = openpyxl.load_workbook(arquivo)
            ws = wb.active
            headers = [str(cell.value or '').strip().lower() for cell in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                d = {headers[i]: str(row[i] or '') for i in range(len(headers)) if i < len(row)}
                clientes.append({
                    'nome': d.get('nome', ''),
                    'cep': d.get('cep', d.get('zip', '')),
                    'rua': d.get('rua', d.get('logradouro', '')),
                    'bairro': d.get('bairro', d.get('neighborhood', '')),
                    'cidade': d.get('cidade', d.get('city', '')),
                    'entrega_retirada': d.get('entrega_retirada', d.get('tipo', 'Entrega')),
                    'observacoes': d.get('observacoes', ''),
                })
        else:
            return jsonify({"erro": "Formato nao suportado. Use CSV, JSON ou XLSX."}), 400
    except Exception as e:
        # `log` ja existe no modulo. O `from ..app import log` que estava aqui
        # subia alem do pacote raiz e levantava ImportError dentro do except,
        # transformando um 400 legivel em 500.
        log.exception("Erro ao ler arquivo importado: %s", e)
        return jsonify({"erro": "Nao foi possivel ler o arquivo enviado. Verifique o formato."}), 400

    if not clientes:
        return jsonify({"erro": "Nenhum cliente encontrado no arquivo."}), 400

    if len(clientes) > MAX_LINHAS_IMPORTACAO:
        return jsonify({
            "erro": f"Arquivo com {len(clientes)} linhas excede o limite de "
                    f"{MAX_LINHAS_IMPORTACAO} por importacao. Divida em partes.",
        }), 400

    clientes_validados = []
    erros_validacao = []
    for idx, c in enumerate(clientes):
        dados_validados, erros = validators.validar_cliente(c)
        if erros:
            erros_validacao.append(f"{c.get('nome', f'linha {idx+1}')}: {'; '.join(erros)}")
        else:
            clientes_validados.append(dados_validados)

    if not clientes_validados:
        return jsonify({"erro": "Nenhum cliente valido para importar.", "detalhes": erros_validacao}), 400

    resultado = database.importar_clientes_em_lote(clientes_validados)
    resposta = {"mensagem": "Importacao concluida!", "sucessos": resultado["sucessos"],
        "duplicados": resultado["duplicados"], "erros": resultado["erros"], "total": len(clientes)}
    if erros_validacao:
        resposta["detalhes_validacao"] = erros_validacao
    return jsonify(resposta)


@clientes_bp.route('/mapa')
def clientes_mapa():
    """Retorna clientes individuais com CEP para geolocalizacao no mapa."""
    return jsonify(database.obter_clientes_para_mapa())


@clientes_bp.route('/backup/json')
def backup_json():
    clientes = database.listar_clientes()
    dados = {"versao": "1.0", "data_exportacao": time.strftime('%Y-%m-%dT%H:%M:%S'), "clientes": clientes}
    buffer = io.BytesIO()
    buffer.write(json.dumps(dados, ensure_ascii=False, indent=2).encode('utf-8'))
    buffer.seek(0)
    return send_file(buffer, mimetype='application/json', as_attachment=True,
        download_name=f"backup_adri_poltronas_{time.strftime('%Y%m%d_%H%M%S')}.json")


@clientes_bp.route('/backup/sqlite')
def backup_sqlite():
    """
    Alias historico de `/api/backup/sqlite`.

    Delega em vez de duplicar: a versao anterior enviava o ARQUIVO EM USO pelo
    `send_file`, o que com o app aberto captura um estado no meio de uma
    transacao e deixa de fora o que estiver no `-wal`. O operador baixava algo
    que parecia um backup e podia nao abrir.
    """
    from blueprints.backup import baixar_sqlite
    return baixar_sqlite()