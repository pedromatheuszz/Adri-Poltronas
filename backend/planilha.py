"""
Construção de planilhas Excel formatadas.

Centraliza o estilo (cabeçalho azul, bordas, largura automática) que antes
estava duplicado entre o envio para o OneDrive e a exportação de relatórios.
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

COR_CABECALHO = '2563EB'
COR_BORDA = 'C0C0C0'
LARGURA_MAX_CONTEUDO = 40
LARGURA_FOLGA = 4

_FONTE_CABECALHO = Font(bold=True, color='FFFFFF', size=11)
_FONTE_CORPO = Font(size=10)
_PREENCHIMENTO_CABECALHO = PatternFill(
    start_color=COR_CABECALHO, end_color=COR_CABECALHO, fill_type='solid'
)
_ALINHAMENTO_CENTRAL = Alignment(horizontal='center', vertical='center')
_LADO_FINO = Side(style='thin', color=COR_BORDA)
_BORDA_FINA = Border(left=_LADO_FINO, right=_LADO_FINO, top=_LADO_FINO, bottom=_LADO_FINO)


def _ajustar_larguras(ws, cabecalhos, total_linhas):
    """Dimensiona cada coluna pelo maior conteúdo, com teto para não estourar."""
    for indice, titulo in enumerate(cabecalhos, start=1):
        largura = len(str(titulo))
        for linha in range(2, total_linhas + 2):
            valor = ws.cell(row=linha, column=indice).value
            if valor:
                largura = max(largura, min(len(str(valor)), LARGURA_MAX_CONTEUDO))
        coluna = openpyxl.utils.get_column_letter(indice)
        ws.column_dimensions[coluna].width = largura + LARGURA_FOLGA


def construir_planilha(titulo, cabecalhos, campos, registros) -> io.BytesIO:
    """
    Monta uma planilha formatada e devolve um buffer posicionado no início.

    titulo     -- nome da aba
    cabecalhos -- rótulos exibidos na primeira linha
    campos     -- chaves de cada registro, na mesma ordem dos cabeçalhos
    registros  -- iterável de dicts
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo

    for indice, rotulo in enumerate(cabecalhos, start=1):
        celula = ws.cell(row=1, column=indice, value=rotulo)
        celula.font = _FONTE_CABECALHO
        celula.fill = _PREENCHIMENTO_CABECALHO
        celula.alignment = _ALINHAMENTO_CENTRAL
        celula.border = _BORDA_FINA

    registros = list(registros)
    for numero_linha, registro in enumerate(registros, start=2):
        for indice, campo in enumerate(campos, start=1):
            celula = ws.cell(row=numero_linha, column=indice, value=registro.get(campo, ''))
            celula.font = _FONTE_CORPO
            celula.alignment = _ALINHAMENTO_CENTRAL
            celula.border = _BORDA_FINA

    _ajustar_larguras(ws, cabecalhos, len(registros))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
